"""
Script de Preenchimento da Planilha Oficial e Exportação PDF da Enel Rio de Janeiro
Formulario_Rateio_ENEL_Rj.xlsm
"""
import os
import sys
import json
import openpyxl

def fill_enel_form(json_data_path, output_dir=None):
    with open(json_data_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    generator_name = data.get('client_name', 'Cliente').strip()
    generator_uc = data.get('utility_id', '').strip()
    generator_doc = data.get('client_document', '').strip()
    gd_type = data.get('gd_type', 'Autoconsumo remoto').strip()
    beneficiaries = data.get('beneficiaries', [])

    clean_name = "".join([c if c.isalnum() or c in (' ', '_', '-') else '_' for c in generator_name]).strip()
    pdf_filename = f"{clean_name}_Formulário_Rateio_ENEL_Rj.pdf"
    xlsm_filename = f"{clean_name}_Formulário_Rateio_ENEL_Rj.xlsm"

    template_path = r"H:\Meu Drive\Trabalhos\Energia fotovoltaica\Alexandre\Formulario_Rateio_ENEL_Rj.xlsm"
    if not os.path.exists(template_path):
        # Fallback local se não estiver na unidade H:
        template_path = os.path.join(os.path.dirname(__file__), "..", "..", "templates", "Formulario_Rateio_ENEL_Rj.xlsm")

    if not output_dir:
        output_dir = os.path.join(os.path.expanduser("~"), "Desktop")

    out_xlsm_path = os.path.join(output_dir, xlsm_filename)
    out_pdf_path = os.path.join(output_dir, pdf_filename)

    # Preencher via openpyxl preservando macros (keep_vba=True)
    if os.path.exists(template_path):
        wb = openpyxl.load_workbook(template_path, keep_vba=True)
        ws = wb[wb.sheetnames[0]]

        # Preenche Unidade Geradora
        ws['C10'] = generator_uc
        ws['G10'] = generator_doc
        ws['M10'] = gd_type

        # Preenche Beneficiárias
        for idx, b in enumerate(beneficiaries):
            row = 23 + idx
            if row > 45:
                break
            ws[f'B{row}'] = b.get('utilityId', '')
            ws[f'C{row}'] = b.get('document', generator_doc)
            pct = float(b.get('percentage', 0)) / 100.0
            ws[f'D{row}'] = pct

        wb.save(out_xlsm_path)
        print(f"EXCEL_SAVED: {out_xlsm_path}")

    # Tenta exportar PDF nativo via Excel COM (se disponível no Windows)
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        
        abs_xlsm = os.path.abspath(out_xlsm_path if os.path.exists(out_xlsm_path) else template_path)
        abs_pdf = os.path.abspath(out_pdf_path)

        wb_com = excel.Workbooks.Open(abs_xlsm)
        # 0 = xlTypePDF
        wb_com.ActiveSheet.ExportAsFixedFormat(0, abs_pdf)
        wb_com.Close(False)
        excel.Quit()
        print(f"PDF_SAVED: {abs_pdf}")
    except Exception as e:
        print(f"COM_EXPORT_NOTE: {e}", file=sys.stderr)

    return out_pdf_path

if __name__ == "__main__":
    if len(sys.argv) > 1:
        json_file = sys.argv[1]
        out_directory = sys.argv[2] if len(sys.argv) > 2 else None
        fill_enel_form(json_file, out_directory)
